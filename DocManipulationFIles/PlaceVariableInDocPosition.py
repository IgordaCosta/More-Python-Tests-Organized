import docx
from docx import Document
import pprint

import FindDocTablePositionsALL




filename='C:\\Users\\IgorDC\\Desktop\\PydocTest\\TestDocNoChange.docx'


splitter='XXXXXXXXXXX'

filename2='C:\\Users\\IgorDC\\Desktop\\PydocTest\\TestChanged.docx'


filename4='C:\\Users\\IgorDC\\Desktop\\PydocTest\\TestChanged4.docx'










PositionList = FindDocTablePositionsALL.FindDocTablePositionsALL(filename)


pprint.pprint(PositionList)

print(len(PositionList))


###########################################
# temp value for test 
# this list will be provided on function
ValueList=[]
for i in range(len(PositionList)):
    ValueList.append("Value_"+str(i+1))
###########################################


doc = docx.Document(filename)
 
# num = 1
# while num <= 10:
for i in range(len(PositionList)):
    [t, r, c ] = PositionList[i]
    doc.tables[t].cell(r,c).text = ValueList[i]

 
doc.save(filename4)