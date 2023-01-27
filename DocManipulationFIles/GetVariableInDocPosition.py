import docx
from docx import Document
import pprint

import FindDocTablePositionsALL




filename='C:\\Users\\IgorDC\\Desktop\\PydocTest\\TestDocNoChange.docx'


splitter='XXXXXXXXXXX'

filename2='C:\\Users\\IgorDC\\Desktop\\PydocTest\\TestChanged.docx'


filename4='C:\\Users\\IgorDC\\Desktop\\PydocTest\\TestChanged4.docx'










PositionList = FindDocTablePositionsALL.FindDocTablePositionsALL(filename)


# pprint.pprint(PositionList)

# print(len(PositionList))




doc = docx.Document(filename4)
 

ValueList=[]
for i in range(len(PositionList)):
    [t, r, c ] = PositionList[i]
    ValueListInput = doc.tables[t].cell(r,c).text
    ValueList.append(ValueListInput)

pprint.pprint(ValueList) 

# print(len(ValueList))
