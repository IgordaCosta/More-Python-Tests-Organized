import docx
from docx import Document
import pprint

import FindDocTablePositionsALL



def GetDocPositionWithVariable(filename,PositionList,LocationKey):
    doc = docx.Document(filename)
    
    DefinedPositionList=[]
    for i in range(len(PositionList)):
        [t, r, c ] = PositionList[i]
        ValueListInput = doc.tables[t].cell(r,c).text
        if ValueListInput==LocationKey:
            DefinedPositionList.append(PositionList[i])

    return DefinedPositionList




LocationKey='XXXXXXXXXXX'

filename='C:\\Users\\IgorDC\\Desktop\\PydocTest\\TestDocNoChange.docx'

filename2='C:\\Users\\IgorDC\\Desktop\\PydocTest\\TestChanged.docx'

filename4='C:\\Users\\IgorDC\\Desktop\\PydocTest\\TestChanged4.docx'


PositionList = FindDocTablePositionsALL.FindDocTablePositionsALL(filename)

DefinedPositionList = GetDocPositionWithVariable(filename2,PositionList,LocationKey)

pprint.pprint(DefinedPositionList) 

print(len(DefinedPositionList))
